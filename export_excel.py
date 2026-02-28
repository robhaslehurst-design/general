"""
Excel Export Module
===================
Exports the full financial model, scenarios, and Monte Carlo results
to a formatted Excel workbook.
"""

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from typing import Dict, List

from financial_model import ModelAssumptions, YearResult, run_model
from scenarios import define_scenarios, run_all_scenarios
from monte_carlo import run_monte_carlo, compute_percentiles


# --- Styling ---
HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SECTION_FONT = Font(name="Calibri", bold=True, size=11, color="2F5496")
SECTION_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
NUM_FONT = Font(name="Calibri", size=10)
MONEY_FMT = '#,##0.0'
PCT_FMT = '0.0%'
PRICE_FMT = '$#,##0.00'
INT_FMT = '#,##0'
THIN_BORDER = Border(
    bottom=Side(style='thin', color='CCCCCC'),
)


def style_header_row(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')


def style_section_row(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = SECTION_FONT
        cell.fill = SECTION_FILL


def write_value(ws, row, col, value, fmt=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = NUM_FONT
    cell.alignment = Alignment(horizontal='right')
    cell.border = THIN_BORDER
    if fmt:
        cell.number_format = fmt


def write_label(ws, row, col, value, indent=0):
    cell = ws.cell(row=row, column=col, value=("  " * indent) + value)
    cell.font = NUM_FONT
    cell.border = THIN_BORDER


def auto_width(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 3, 25)


# ---------------------------------------------------------------------------
# Sheet 1: Assumptions
# ---------------------------------------------------------------------------
def write_assumptions_sheet(wb: Workbook, assumptions: ModelAssumptions):
    ws = wb.create_sheet("Assumptions")

    params = [
        ("REVENUE ASSUMPTIONS", None, None),
        ("Number of Partners (Year 0)", assumptions.num_partners, INT_FMT),
        ("Revenue per Partner ($M)", assumptions.revenue_per_partner_m, MONEY_FMT),
        ("Partner Growth Rate (p.a.)", assumptions.partner_growth_rate, PCT_FMT),
        ("RPP Target Year 5 ($M)", assumptions.rpp_target_m, MONEY_FMT),
        ("Consulting Staff (Year 0)", assumptions.num_consulting_staff, INT_FMT),
        ("", None, None),
        ("MARGIN ASSUMPTIONS", None, None),
        ("Gross Margin (Year 0)", assumptions.gross_margin_pct, PCT_FMT),
        ("Gross Margin Target (Year 5)", assumptions.gross_margin_target_pct, PCT_FMT),
        ("Operating Margin (Year 0)", assumptions.operating_margin_pct, PCT_FMT),
        ("Operating Margin Target (Year 5)", assumptions.operating_margin_target_pct, PCT_FMT),
        ("", None, None),
        ("COST STRUCTURE", None, None),
        ("Partner Base Pay ($K)", assumptions.partner_base_pay_k, '#,##0'),
        ("Investment (% of Revenue)", assumptions.investment_pct_revenue, PCT_FMT),
        ("", None, None),
        ("BUSINESS MIX", None, None),
        ("Market Studies (%)", assumptions.market_studies_pct, PCT_FMT),
        ("Transformation / Strategy (%)", assumptions.transformation_strategy_pct, PCT_FMT),
        ("", None, None),
        ("EBITDA & VALUATION", None, None),
        ("EBITDA Fixed (% of Revenue)", assumptions.ebitda_fixed_pct, PCT_FMT),
        ("EBITDA Residual Share (1/3)", assumptions.ebitda_residual_share, '0.00'),
        ("Valuation Multiple (x EBITDA)", assumptions.valuation_multiple, '0.00x'),
        ("Structural Debt ($M)", assumptions.structural_debt_m, MONEY_FMT),
        ("Total Shares", assumptions.total_shares, '#,##0'),
        ("", None, None),
        ("CAPITAL & TAX", None, None),
        ("Working Capital Days", assumptions.working_capital_days, INT_FMT),
        ("Dividend Yield", assumptions.dividend_yield, PCT_FMT),
        ("Borrowing Rate", assumptions.borrowing_rate, PCT_FMT),
        ("Corporate Tax Rate (US)", assumptions.tax_rate, PCT_FMT),
    ]

    # Headers
    ws.cell(row=1, column=1, value="Parameter").font = HEADER_FONT
    ws.cell(row=1, column=2, value="Value").font = HEADER_FONT
    ws.cell(row=1, column=1).fill = HEADER_FILL
    ws.cell(row=1, column=2).fill = HEADER_FILL

    for i, (name, value, fmt) in enumerate(params, start=2):
        if value is None and name:
            ws.cell(row=i, column=1, value=name).font = SECTION_FONT
            ws.cell(row=i, column=1).fill = SECTION_FILL
            ws.cell(row=i, column=2).fill = SECTION_FILL
        elif name:
            write_label(ws, i, 1, name)
            if value is not None:
                write_value(ws, i, 2, value, fmt)

    auto_width(ws)


# ---------------------------------------------------------------------------
# Sheet 2: Base Case P&L
# ---------------------------------------------------------------------------
def write_pnl_sheet(wb: Workbook, results: List[YearResult], sheet_name: str = "Base Case"):
    ws = wb.create_sheet(sheet_name)

    n_years = len(results)
    max_col = n_years + 1

    # Header row
    ws.cell(row=1, column=1, value="Financial Forecast ($M)")
    for j, r in enumerate(results):
        ws.cell(row=1, column=j+2, value=f"Year {r.year}")
    style_header_row(ws, 1, max_col)

    # Define rows: (label, getter, format, indent, is_section)
    rows = [
        ("REVENUE BUILD-UP", None, None, 0, True),
        ("Partners (#)", lambda r: r.num_partners, INT_FMT, 0, False),
        ("Revenue per Partner ($M)", lambda r: r.revenue_per_partner_m, MONEY_FMT, 0, False),
        ("Total Revenue", lambda r: r.total_revenue_m, MONEY_FMT, 0, False),
        ("Market Studies Revenue", lambda r: r.market_studies_revenue_m, MONEY_FMT, 1, False),
        ("Transformation / Strategy", lambda r: r.transformation_revenue_m, MONEY_FMT, 1, False),
        ("", None, None, 0, False),
        ("COST OF SALES", None, None, 0, True),
        ("Partner Base Cost", lambda r: r.partner_base_cost_m, MONEY_FMT, 1, False),
        ("Consulting Staff Cost", lambda r: r.consulting_staff_cost_m, MONEY_FMT, 1, False),
        ("Other Cost of Sales", lambda r: r.other_cos_m, MONEY_FMT, 1, False),
        ("Total Cost of Sales", lambda r: r.total_cos_m, MONEY_FMT, 0, False),
        ("Gross Profit", lambda r: r.gross_profit_m, MONEY_FMT, 0, False),
        ("Gross Margin", lambda r: r.gross_margin_pct, PCT_FMT, 0, False),
        ("", None, None, 0, False),
        ("OPERATING EXPENSES", None, None, 0, True),
        ("Total Operating Expenses", lambda r: r.total_opex_m, MONEY_FMT, 0, False),
        ("Operating Profit", lambda r: r.operating_profit_m, MONEY_FMT, 0, False),
        ("Operating Margin", lambda r: r.operating_margin_pct, PCT_FMT, 0, False),
        ("", None, None, 0, False),
        ("EBITDA & PARTNER ECONOMICS", None, None, 0, True),
        ("EBITDA", lambda r: r.ebitda_m, MONEY_FMT, 0, False),
        ("Partner Bonus Pool", lambda r: r.partner_bonus_pool_m, MONEY_FMT, 1, False),
        ("Total Partner Compensation", lambda r: r.total_partner_comp_m, MONEY_FMT, 0, False),
        ("Per-Partner Total Comp ($M)", lambda r: r.per_partner_total_comp_m, MONEY_FMT, 0, False),
        ("", None, None, 0, False),
        ("BELOW THE LINE", None, None, 0, True),
        ("Depreciation", lambda r: r.depreciation_m, MONEY_FMT, 1, False),
        ("Interest Expense", lambda r: r.interest_expense_m, MONEY_FMT, 1, False),
        ("Tax", lambda r: r.tax_m, MONEY_FMT, 1, False),
        ("Net Income", lambda r: r.net_income_m, MONEY_FMT, 0, False),
        ("", None, None, 0, False),
        ("VALUATION", None, None, 0, True),
        ("Enterprise Value", lambda r: r.enterprise_value_m, MONEY_FMT, 0, False),
        ("Equity Value", lambda r: r.equity_value_m, MONEY_FMT, 0, False),
        ("Share Price ($)", lambda r: r.share_price, PRICE_FMT, 0, False),
        ("Dividend per Share ($)", lambda r: r.dividend_per_share, PRICE_FMT, 0, False),
        ("Total Dividends", lambda r: r.total_dividends_m, MONEY_FMT, 0, False),
        ("", None, None, 0, False),
        ("WORKING CAPITAL & STAFF", None, None, 0, True),
        ("Working Capital", lambda r: r.working_capital_m, MONEY_FMT, 0, False),
        ("Capex (Investment)", lambda r: r.capex_m, MONEY_FMT, 0, False),
        ("Consulting Staff (#)", lambda r: r.num_consulting_staff, INT_FMT, 0, False),
        ("Revenue per Consultant ($K)", lambda r: r.revenue_per_consultant_k, '#,##0', 0, False),
        ("Leverage (Staff/Partner)", lambda r: r.leverage_ratio, '0.0', 0, False),
    ]

    row_num = 2
    for label, getter, fmt, indent, is_section in rows:
        if not label:
            row_num += 1
            continue
        if is_section:
            ws.cell(row=row_num, column=1, value=label)
            style_section_row(ws, row_num, max_col)
            row_num += 1
            continue

        write_label(ws, row_num, 1, label, indent)
        for j, r in enumerate(results):
            write_value(ws, row_num, j+2, getter(r), fmt)
        row_num += 1

    auto_width(ws)


# ---------------------------------------------------------------------------
# Sheet 3: Scenario Comparison
# ---------------------------------------------------------------------------
def write_scenario_comparison(wb: Workbook, all_results: Dict[str, List[YearResult]]):
    ws = wb.create_sheet("Scenario Comparison")

    scenarios = list(all_results.keys())
    max_col = len(scenarios) + 1

    ws.cell(row=1, column=1, value="Year 5 Comparison")
    for j, s in enumerate(scenarios):
        ws.cell(row=1, column=j+2, value=s)
    style_header_row(ws, 1, max_col)

    metrics = [
        ("Revenue ($M)", lambda r: r.total_revenue_m, MONEY_FMT),
        ("Gross Margin", lambda r: r.gross_margin_pct, PCT_FMT),
        ("Operating Profit ($M)", lambda r: r.operating_profit_m, MONEY_FMT),
        ("Operating Margin", lambda r: r.operating_margin_pct, PCT_FMT),
        ("EBITDA ($M)", lambda r: r.ebitda_m, MONEY_FMT),
        ("Net Income ($M)", lambda r: r.net_income_m, MONEY_FMT),
        ("Share Price ($)", lambda r: r.share_price, PRICE_FMT),
        ("Dividend/Share ($)", lambda r: r.dividend_per_share, PRICE_FMT),
        ("Partners (#)", lambda r: r.num_partners, INT_FMT),
        ("Per-Partner Comp ($M)", lambda r: r.per_partner_total_comp_m, MONEY_FMT),
        ("Working Capital ($M)", lambda r: r.working_capital_m, MONEY_FMT),
    ]

    for i, (label, getter, fmt) in enumerate(metrics, start=2):
        write_label(ws, i, 1, label)
        for j, s in enumerate(scenarios):
            yr5 = all_results[s][-1]
            write_value(ws, i, j+2, getter(yr5), fmt)

    auto_width(ws)


# ---------------------------------------------------------------------------
# Sheet 4: Monte Carlo Percentiles
# ---------------------------------------------------------------------------
def write_monte_carlo_sheet(wb: Workbook, arrays: Dict[str, np.ndarray], n_sims: int):
    ws = wb.create_sheet("Monte Carlo")

    pcts_data = compute_percentiles(arrays)
    n_years = arrays["revenue"].shape[1]

    key_metrics = [
        ("Revenue ($M)", "revenue", MONEY_FMT),
        ("EBITDA ($M)", "ebitda", MONEY_FMT),
        ("Operating Margin", "operating_margin", PCT_FMT),
        ("Share Price ($)", "share_price", PRICE_FMT),
        ("Dividend/Share ($)", "dividend_per_share", PRICE_FMT),
        ("Per-Partner Comp ($M)", "per_partner_comp", MONEY_FMT),
        ("Equity Value ($M)", "equity_value", MONEY_FMT),
        ("Net Income ($M)", "net_income", MONEY_FMT),
    ]

    ws.cell(row=1, column=1, value=f"Monte Carlo — {n_sims:,} Simulations")
    ws.cell(row=1, column=1).font = Font(name="Calibri", bold=True, size=12, color="2F5496")

    row_num = 3
    for label, metric, fmt in key_metrics:
        ws.cell(row=row_num, column=1, value=label)
        style_section_row(ws, row_num, n_years + 2)
        row_num += 1

        # Sub-header
        ws.cell(row=row_num, column=1, value="Percentile")
        for y in range(n_years):
            ws.cell(row=row_num, column=y+2, value=f"Year {y}")
        style_header_row(ws, row_num, n_years + 1)
        row_num += 1

        for p in [5, 25, 50, 75, 95]:
            write_label(ws, row_num, 1, f"P{p}")
            for y in range(n_years):
                write_value(ws, row_num, y+2, float(pcts_data[metric][p][y]), fmt)
            row_num += 1

        # Mean row
        write_label(ws, row_num, 1, "Mean")
        means = np.mean(arrays[metric], axis=0)
        for y in range(n_years):
            write_value(ws, row_num, y+2, float(means[y]), fmt)
        row_num += 1

        # Std Dev row
        write_label(ws, row_num, 1, "Std Dev")
        stds = np.std(arrays[metric], axis=0)
        for y in range(n_years):
            write_value(ws, row_num, y+2, float(stds[y]), fmt)
        row_num += 1

        row_num += 1  # blank row

    auto_width(ws)


# ---------------------------------------------------------------------------
# Main export function
# ---------------------------------------------------------------------------
def export_to_excel(filename: str = "consulting_firm_forecast.xlsx", n_mc_sims: int = 5000):
    """Generate the complete Excel workbook."""
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Assumptions
    base = ModelAssumptions()
    write_assumptions_sheet(wb, base)

    # 2. Base Case P&L
    base_results = run_model(base)
    write_pnl_sheet(wb, base_results, "Base Case")

    # 3. All scenario P&Ls
    all_results = run_all_scenarios()
    for name, results in all_results.items():
        if name != "Base Case":
            write_pnl_sheet(wb, results, name)

    # 4. Scenario comparison
    write_scenario_comparison(wb, all_results)

    # 5. Monte Carlo
    print(f"  Running {n_mc_sims:,} Monte Carlo simulations...")
    arrays, _ = run_monte_carlo(n_simulations=n_mc_sims)
    write_monte_carlo_sheet(wb, arrays, n_mc_sims)

    wb.save(filename)
    print(f"  Saved to {filename}")
    return filename


if __name__ == "__main__":
    export_to_excel()
